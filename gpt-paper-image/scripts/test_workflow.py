"""Synthetic-only regression checks; no research documents or image fixtures."""
import hashlib
from pathlib import Path
import tempfile
import unittest

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

import bundle
from audit_pptx import audit
from native_shapes import clear_effects, linear_gradient


def fixture(path, picture=False, group=False):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shapes = slide.shapes.add_group_shape().shapes if group else slide.shapes
    box = shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(1), Inches(3), Inches(1))
    box.text = 'Input to output'
    linear_gradient(box, 'FFFFFF', 'DDEEFF')
    clear_effects(box)
    if picture:
        image = path.parent / 'synthetic-pixel.png'
        Image.new('RGB', (4, 4), 'white').save(image)
        slide.shapes.add_picture(str(image), Inches(6), Inches(1), Inches(1), Inches(1))
    prs.save(path)


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='gpt-paper-image-test-')
        self.base = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def test_preserve_original_and_reject_overwrite(self):
        root = bundle.initialize(self.base / 'run')
        original = self.base / 'synthetic.png'
        Image.new('RGB', (4, 4), 'white').save(original)
        item = bundle.add(root, original, 'generated-original', 'generated/v001.png')
        self.assertEqual(item['sha256'], bundle.digest(original))
        self.assertTrue(bundle.verify(root)['ok'])
        Image.new('RGB', (4, 4), 'black').save(original)
        with self.assertRaises(ValueError):
            bundle.add(root, original, 'generated-original', 'generated/v001.png')
        self.assertEqual(bundle.digest(root / 'generated/v001.png'), item['sha256'])

    def test_missing_roles_and_modified_file(self):
        root = bundle.initialize(self.base / 'run')
        text = self.base / 'prompt.txt'; text.write_text('A generic diagram.')
        bundle.add(root, text, 'prompt', 'prompts/v001.txt')
        self.assertFalse(bundle.verify(root, complete=True)['ok'])
        (root / 'prompts/v001.txt').write_text('Changed after registration.')
        self.assertIn('Hash mismatch: prompts/v001.txt', bundle.verify(root)['errors'])

    def test_path_and_skill_boundary(self):
        root = bundle.initialize(self.base / 'run')
        with self.assertRaises(ValueError):
            bundle.member(root, '../outside.txt')
        (root / 'escape').symlink_to(self.base, target_is_directory=True)
        with self.assertRaises(ValueError):
            bundle.member(root, 'escape/outside.txt')
        with self.assertRaises(ValueError):
            bundle.root_path(bundle.SKILL_ROOT / 'private-artifacts')

    def test_handoff_ready_without_generated_image_or_ppt(self):
        root = bundle.initialize(self.base / 'run')
        self.assertTrue((root / 'references/images').is_dir())
        self.assertTrue((root / 'handoff').is_dir())
        image = self.base / 'reference.png'
        Image.new('RGB', (4, 4), 'white').save(image)
        notes = self.base / 'notes.json'; notes.write_text('{}')
        prompt = self.base / 'prompt.txt'; prompt.write_text('A generic diagram.')
        handoff = self.base / 'handoff.md'; handoff.write_text('Awaiting user image.')
        for src, role, rel in [
            (image, 'reference', 'references/images/ref-01.png'),
            (notes, 'reference-notes', 'references/style-references.json'),
            (prompt, 'prompt', 'prompts/v001.txt'),
            (handoff, 'handoff', 'handoff/web-generation.md'),
        ]:
            bundle.add(root, src, role, rel)
        self.assertTrue(bundle.verify(root, stage='handoff')['ok'])
        self.assertFalse(bundle.verify(root, complete=True)['ok'])
        (root / 'references/images/ref-01.png').write_bytes(b'changed')
        self.assertIn('Hash mismatch: references/images/ref-01.png',
                      bundle.verify(root, stage='handoff')['errors'])

    def test_handoff_requires_reference_image_and_instructions(self):
        root = bundle.initialize(self.base / 'run')
        text = self.base / 'prompt.txt'; text.write_text('A generic diagram.')
        bundle.add(root, text, 'prompt', 'prompts/v001.txt')
        result = bundle.verify(root, stage='handoff')
        self.assertFalse(result['ok'])
        self.assertIn('Missing role: reference', result['errors'])
        self.assertIn('Missing role: handoff', result['errors'])
        with self.assertRaises(ValueError):
            bundle.verify(root, complete=True, stage='handoff')

    def test_native_gradient_roundtrip_without_changing_original(self):
        path = self.base / 'native.pptx'; fixture(path)
        sha = hashlib.sha256(path.read_bytes()).hexdigest()
        result = audit(path)
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['checks']['gradient_roundtrip'], 'PASS')
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(), sha)

    def test_grouped_native_objects_are_checked(self):
        path = self.base / 'grouped.pptx'; fixture(path, group=True)
        result = audit(path)
        self.assertTrue(result['ok'], result)
        self.assertEqual(result['slides'][0]['gradients'], 1)

    def test_picture_is_not_reported_as_all_native(self):
        path = self.base / 'with-picture.pptx'; fixture(path, picture=True)
        result = audit(path)
        self.assertFalse(result['ok'])
        self.assertEqual(result['slides'][0]['picture_objects'], 1)
        self.assertEqual(result['slides'][0]['image_fills'], 0)
        permitted = audit(path, allow_media=True)
        self.assertTrue(permitted['ok'])
        self.assertFalse(permitted['all_native'])
        self.assertTrue(permitted['media_files'])

    def test_image_filled_shape_is_not_missed(self):
        path = self.base / 'image-fill.pptx'; fixture(path)
        image = self.base / 'fill.png'; Image.new('RGB', (4, 4), 'white').save(image)
        prs = Presentation(path); slide = prs.slides[0]
        _, rid = slide.part.get_or_add_image_part(str(image))
        sp = slide.shapes[0]._element.spPr
        for child in list(sp):
            if child.tag.split('}')[-1] == 'gradFill':sp.remove(child)
        fill = OxmlElement('a:blipFill'); blip = OxmlElement('a:blip')
        blip.set(qn('r:embed'), rid); fill.append(blip)
        stretch = OxmlElement('a:stretch'); stretch.append(OxmlElement('a:fillRect')); fill.append(stretch)
        sp.insert(2, fill); prs.save(path)
        result = audit(path)
        self.assertFalse(result['ok'])
        self.assertEqual(result['slides'][0]['picture_objects'], 0)
        self.assertEqual(result['slides'][0]['image_fills'], 1)


if __name__ == '__main__':
    unittest.main(verbosity=2)

#!/usr/bin/env python3
"""Inspect native objects and perform a temporary edit/save/reload check."""
import argparse
import json
from pathlib import Path
import sys
import tempfile
import zipfile


def walk(shapes, prefix=()):
    for i, shape in enumerate(shapes):
        path = prefix + (i,)
        yield path, shape
        if hasattr(shape, 'shapes'):
            yield from walk(shape.shapes, path)


def locate(slide, indices):
    shapes = slide.shapes
    for idx in indices:
        shape = shapes[idx]
        if hasattr(shape, 'shapes'):
            shapes = shape.shapes
    return shape


def audit(filename, pdf=None, allow_media=False, minimum_pt=5.0):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    source = Path(filename).resolve(strict=True)
    prs = Presentation(source)
    errors = []
    report = {'slides': [], 'checks': {}, 'scope':
        'Structural/editability checks only; visual fidelity and scientific meaning require review.'}
    if not prs.slides:
        errors.append('The deck has no slides.')
    with zipfile.ZipFile(source) as archive:
        media = [n for n in archive.namelist() if n.startswith('ppt/media/')]
    report['media_files'] = media
    if media and not allow_media:
        errors.append('Embedded media is present in an all-native deliverable.')
    text_target = gradient_target = None
    for si, slide in enumerate(prs.slides):
        items = list(walk(slide.shapes))
        pictures = sum(s.shape_type == MSO_SHAPE_TYPE.PICTURE for _, s in items)
        # Ordinary p:pic objects are counted above; image_fills means an image
        # used as the fill of a native shape, not that same picture counted twice.
        blips = sum(len(s._element.xpath('./p:spPr/a:blipFill/a:blip')) for _, s in items)
        gradients = sum(bool(s._element.xpath('./p:spPr/a:gradFill')) for _, s in items)
        texts = [(path, s) for path, s in items if s.has_text_frame and s.text.strip()]
        report['slides'].append({'objects_including_groups': len(items), 'text_objects': len(texts),
                                  'picture_objects': pictures, 'image_fills': blips, 'gradients': gradients})
        if (pictures or blips) and not allow_media:
            errors.append(f'Slide {si+1} has picture objects or image fills.')
        if not texts:
            errors.append(f'Slide {si+1} has no editable text; review the conversion.')
        if texts and text_target is None:
            text_target = (si, texts[0][0])
        for path, shape in items:
            if shape._element.xpath('./p:spPr/a:gradFill/a:gsLst/a:gs/a:srgbClr') and gradient_target is None:
                gradient_target = (si, path)
    report['all_native'] = not media and not any(s['picture_objects'] or s['image_fills'] for s in report['slides'])
    if text_target is not None:
        si, indices = text_target
        shape = locate(prs.slides[si], indices)
        shape.text = 'Editability check'
        shape.left += 9144
        new_left = shape.left
        if gradient_target is not None:
            gs, gp = gradient_target
            color = locate(prs.slides[gs], gp)._element.xpath('./p:spPr/a:gradFill/a:gsLst/a:gs/a:srgbClr')[0]
            color.set('val', 'ABCDEF')
        with tempfile.TemporaryDirectory(prefix='ppt-editability-') as temp:
            copy = Path(temp) / 'roundtrip.pptx'
            prs.save(copy)
            again = Presentation(copy)
            changed = locate(again.slides[si], indices)
            assert changed.text == 'Editability check' and changed.left == new_left
            if gradient_target is not None:
                color = locate(again.slides[gs], gp)._element.xpath('./p:spPr/a:gradFill/a:gsLst/a:gs/a:srgbClr')[0]
                assert color.get('val') == 'ABCDEF'
        report['checks']['text_position_roundtrip'] = 'PASS'
        report['checks']['gradient_roundtrip'] = 'PASS' if gradient_target else 'NOT APPLICABLE'
    else:
        report['checks']['text_position_roundtrip'] = 'NOT TESTED'
    if pdf:
        import fitz
        with fitz.open(pdf) as doc:
            if len(doc) != len(prs.slides):
                errors.append('PDF page count differs from slide count.')
            pages = []
            for i, page in enumerate(doc):
                spans = [s for b in page.get_text('dict')['blocks'] if 'lines' in b
                         for line in b['lines'] for s in line['spans'] if s['text'].strip()]
                smallest = min((s['size'] for s in spans), default=None)
                images = len(page.get_images())
                if smallest is None:
                    errors.append(f'PDF page {i+1} has no extractable text.')
                elif smallest < minimum_pt:
                    errors.append(f'PDF page {i+1} minimum font {smallest:.3f} pt is below {minimum_pt}.')
                pages.append({'page': i+1, 'minimum_font_pt': smallest, 'image_resources': images})
            report['pdf'] = {'pages': pages, 'minimum_required_pt': minimum_pt,
                'vector_only': all(p['image_resources'] == 0 for p in pages),
                'provenance_note': 'The caller must verify this PDF was exported from the final PPT.'}
    report['errors'] = errors
    report['ok'] = not errors
    report['visual_review'] = 'REQUIRED SEPARATELY'
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('pptx'); parser.add_argument('--pdf')
    parser.add_argument('--output'); parser.add_argument('--min-pt', type=float, default=5.0)
    parser.add_argument('--allow-media', action='store_true', help='For explicitly allowed mixed-media figures; media remain reported.')
    args = parser.parse_args()
    try:
        result = audit(args.pptx, args.pdf, args.allow_media, args.min_pt)
        text = json.dumps(result, indent=2, ensure_ascii=False)
        if args.output:
            output = Path(args.output).expanduser().resolve()
            skill = Path(__file__).resolve().parents[1]
            if skill == output or skill in output.parents:
                raise ValueError('Write task reports outside the skill package.')
            if output.exists():
                raise ValueError('QA report already exists; choose a new version path.')
            output.parent.mkdir(parents=True, exist_ok=True)
            output.write_text(text+'\n', encoding='utf-8')
        print(text)
        return 0 if result['ok'] else 1
    except (OSError, ValueError, ImportError, AssertionError) as exc:
        print(f'Audit could not complete: {exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
